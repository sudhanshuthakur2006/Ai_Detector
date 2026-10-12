import io
from docx import Document
from file_processor import FileProcessor
from detector import AIDetector
from fastapi.testclient import TestClient
from main import app

client = TestClient(app)

def test_file_processor_txt():
    sample_text = b"Hello world.\nThis is a student essay paragraph.\n- Bullet point one\n- Bullet point two"
    text, paras = FileProcessor.extract_from_txt(sample_text)
    assert len(paras) == 4
    assert paras[2]["is_bullet"] is True

    reconstructed = FileProcessor.reconstruct_txt(paras)
    assert b"Hello world" in reconstructed
    assert b"Bullet point one" in reconstructed

def test_file_processor_docx():
    # Build sample docx
    doc = Document()
    doc.add_heading("Research Title", level=1)
    p = doc.add_paragraph("This is the introductory paragraph with some thoughts.")
    doc.add_paragraph("First bullet item", style="List Bullet")
    
    bio = io.BytesIO()
    doc.save(bio)
    docx_bytes = bio.getvalue()

    # Extract
    extracted_text, paras = FileProcessor.extract_from_docx(docx_bytes)
    assert "Research Title" in extracted_text
    assert len(paras) == 3

    # Reconstruct
    paras[1]["text"] = "This is the humanized introductory paragraph."
    reconstructed_bytes = FileProcessor.reconstruct_docx(docx_bytes, paras)
    assert len(reconstructed_bytes) > 0

    # Verify reconstructed docx
    new_doc = Document(io.BytesIO(reconstructed_bytes))
    new_texts = [p.text for p in new_doc.paragraphs if p.text.strip()]
    assert "This is the humanized introductory paragraph." in new_texts

def test_detector_scoring():
    detector = AIDetector()
    
    # Typical AI-sounding repetitive monotonous text with buzzwords
    ai_text = (
        "Furthermore, we must delve into the intricate nuances of the dynamic paradigm. "
        "Moreover, this demonstrates a robust framework to harness and leverage cutting-edge capabilities. "
        "In conclusion, the testament to modern innovation underscores multifaceted dimensions."
    )
    ai_result = detector.detect_text(ai_text)
    assert ai_result["ai_percentage"] > 40.0
    assert "human_percentage" in ai_result
    assert "mixed_percentage" in ai_result
    assert sum([ai_result["ai_percentage"], ai_result["human_percentage"], ai_result["mixed_percentage"]]) > 99.0

    # Natural student human text with high burstiness
    human_text = (
        "Honestly? I was pretty skeptical at first. "
        "The experiments took way longer than expected because our initial data pipeline had a dumb bug in the parsing logic, which ended up skewing all the timestamps. "
        "Once we fixed that, things started looking reasonable. Not perfect, but enough to make sense of the results."
    )
    human_result = detector.detect_text(human_text)
    assert human_result["human_percentage"] > 50.0

def test_api_endpoints():
    # Health
    res = client.get("/api/health")
    assert res.status_code == 200
    assert res.json()["status"] == "online"

    # Detect
    res = client.post("/api/detect", json={"text": "This is a quick test sentence."})
    assert res.status_code == 200
    data = res.json()
    assert "ai_percentage" in data

    # Upload txt
    txt_content = b"Introduction to machine learning.\nNatural human writing with diverse sentence structure."
    files = {"file": ("test.txt", txt_content, "text/plain")}
    res = client.post("/api/upload", files=files)
    assert res.status_code == 200
    upload_data = res.json()
    assert "session_id" in upload_data
    assert upload_data["total_paragraphs"] == 2

if __name__ == "__main__":
    print("Running test_file_processor_txt...")
    test_file_processor_txt()
    print("Running test_file_processor_docx...")
    test_file_processor_docx()
    print("Running test_detector_scoring...")
    test_detector_scoring()
    print("Running test_api_endpoints...")
    test_api_endpoints()
    print("\nALL BACKEND TESTS PASSED!")
